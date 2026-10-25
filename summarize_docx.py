"""Create a concise, extractive summary DOCX from an existing Word document."""

from __future__ import annotations

import argparse
import math
import os
import re
import subprocess
import sys
import tempfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph
from docx.table import Table


STOP_WORDS = {
    "about", "above", "after", "again", "against", "also", "among", "and",
    "are", "because", "been", "before", "being", "below", "between", "both",
    "but", "can", "could", "does", "each", "for", "from", "have", "having",
    "into", "its", "more", "most", "not", "only", "other", "over", "same",
    "should", "some", "such", "than", "that", "the", "their", "them", "then",
    "there", "these", "they", "this", "those", "through", "under", "until",
    "was", "were", "what", "when", "where", "which", "while", "with", "would",
    "you", "your",
}
WORD_RE = re.compile(r"\b[a-zA-Z][a-zA-Z'-]*\b")
SENTENCE_RE = re.compile(r"(?<=[.!?])\s+")
TASK_NAME = "AutomationDocxDescriptionDaily"


def install_daily_schedule(script_path: Path) -> None:
    """Register this script to run every day at 10 AM local time."""
    if sys.platform != "win32":
        raise OSError("Daily scheduling is currently supported on Windows only.")

    action = f'"{Path(sys.executable).resolve()}" "{script_path.resolve()}"'
    result = subprocess.run(
        [
            "schtasks",
            "/Create",
            "/SC",
            "DAILY",
            "/ST",
            "10:00",
            "/TN",
            TASK_NAME,
            "/TR",
            action,
            "/F",
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0:
        details = result.stderr.strip() or result.stdout.strip()
        raise OSError(f"Could not create the daily scheduled task: {details}")

    print(f"Scheduled {script_path.name} to run daily at 10:00 AM local time.")
    if result.stdout.strip():
        print(result.stdout.strip())


def iter_document_text(document: object) -> list[str]:
    """Return non-empty body paragraph and table-cell text in document order."""
    text_blocks: list[str] = []
    body = document.element.body

    for block in body.iterchildren():
        if isinstance(block, CT_P):
            text = Paragraph(block, document).text.strip()
            if text:
                text_blocks.append(text)
        elif isinstance(block, CT_Tbl):
            table = Table(block, document)
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        text_blocks.append(text)

    return text_blocks


def summarize(text_blocks: list[str], max_sentences: int = 3) -> str:
    """Select representative sentences using document word frequencies."""
    sentences = [
        sentence.strip()
        for block in text_blocks
        for sentence in SENTENCE_RE.split(block)
        if sentence.strip()
    ]
    if not sentences:
        raise ValueError("The input DOCX contains no extractable text to summarize.")

    words_by_sentence = [
        [word.lower() for word in WORD_RE.findall(sentence)]
        for sentence in sentences
    ]
    frequencies = Counter(
        word
        for words in words_by_sentence
        for word in words
        if len(word) > 2 and word not in STOP_WORDS
    )

    if not frequencies:
        return " ".join(sentences[:max_sentences])

    max_frequency = max(frequencies.values())
    scored_sentences: list[tuple[float, int, str]] = []
    for index, (sentence, words) in enumerate(zip(sentences, words_by_sentence)):
        keywords = [word for word in words if len(word) > 2 and word not in STOP_WORDS]
        if not keywords:
            continue
        score = sum(frequencies[word] / max_frequency for word in keywords)
        score /= math.sqrt(len(keywords))
        if index == 0:
            score *= 1.2
        scored_sentences.append((score, index, sentence))

    if not scored_sentences:
        return " ".join(sentences[:max_sentences])

    opening = next(
        (
            item
            for item in scored_sentences
            if len(words_by_sentence[item[1]]) >= 8
        ),
        None,
    )
    remaining = [
        item for item in scored_sentences if opening is None or item[1] != opening[1]
    ]
    ranked = sorted(remaining, reverse=True)[: max_sentences - (opening is not None)]
    selected = sorted(
        ([opening] if opening is not None else []) + ranked,
        key=lambda item: item[1],
    )
    return " ".join(sentence for _, _, sentence in selected)


def find_input(directory: Path) -> Path:
    candidates = sorted(
        path
        for path in directory.glob("*.docx")
        if not path.name.startswith("~$")
        and not path.stem.endswith("_description")
    )
    if not candidates:
        raise FileNotFoundError(f"No source DOCX files found in {directory}.")
    if len(candidates) > 1:
        names = ", ".join(path.name for path in candidates)
        raise ValueError(
            f"More than one source DOCX was found ({names}). "
            "Specify one with --input."
        )
    return candidates[0]


def save_description(output_path: Path, input_path: Path, summary: str) -> Path:
    output_document = Document()
    output_document.add_heading(f"Description of {input_path.stem}", level=0)
    output_document.add_paragraph(f"Source document: {input_path.name}")
    output_document.add_heading("Summary", level=1)
    output_document.add_paragraph(summary)

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{output_path.stem}_",
            suffix=".docx",
            dir=output_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        output_document.save(temporary_path)
        try:
            os.replace(temporary_path, output_path)
            return output_path
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            fallback_path = output_path.with_name(
                f"{output_path.stem}_{timestamp}{output_path.suffix}"
            )
            os.replace(temporary_path, fallback_path)
            return fallback_path
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def create_description(input_path: Path, output_path: Path) -> Path:
    source_document = Document(input_path)
    content = iter_document_text(source_document)
    summary = summarize(content)
    return save_description(output_path, input_path, summary)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Summarize a DOCX document into a separate DOCX file."
    )
    parser.add_argument(
        "--input",
        type=Path,
        help="Source DOCX. Defaults to the only source DOCX beside this script.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="Destination DOCX. Defaults to <source>_description.docx.",
    )
    parser.add_argument(
        "--install-schedule",
        action="store_true",
        help="Schedule this script to run daily at 10:00 AM local time on Windows.",
    )
    args = parser.parse_args()

    try:
        script_directory = Path(__file__).resolve().parent
        if args.install_schedule:
            install_daily_schedule(Path(__file__))
            return 0

        input_path = (args.input or find_input(script_directory)).resolve()
        if not input_path.is_file():
            raise FileNotFoundError(f"Input file does not exist: {input_path}")
        if input_path.suffix.lower() != ".docx":
            raise ValueError(f"Input must be a .docx file: {input_path}")

        output_path = (
            args.output.resolve()
            if args.output
            else input_path.with_name(f"{input_path.stem}_description.docx")
        )
        if input_path == output_path:
            raise ValueError("Input and output paths must be different.")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        actual_output_path = create_description(input_path, output_path)
    except (FileNotFoundError, ValueError, OSError) as error:
        print(
            f"Error: {error}\n"
            "Check that the output folder is writable. If the destination DOCX "
            "is open in Word, close it and try again.",
            file=sys.stderr,
        )
        return 1

    if actual_output_path != output_path:
        print(
            f"The destination DOCX is locked or unavailable. "
            f"Created a separate copy instead: {actual_output_path}"
        )
    else:
        print(f"Created summary DOCX: {actual_output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
